from docx import Document

def create_master():
    doc = Document()
    
    doc.add_heading('MANIDHAR PATI', 0)
    doc.add_paragraph('Machine Learning Engineer | +91 8008106844 | manidharmanidhar2@gmail.com')
    
    # AI TARGET SECTION
    doc.add_heading('PROFESSIONAL SUMMARY', level=1)
    doc.add_paragraph("{{AI_SUMMARY_INJECTION_POINT}}") # AI will write here
    
    # AI TARGET SECTION
    doc.add_heading('PROFESSIONAL EXPERIENCE', level=1)
    doc.add_paragraph("{{AI_EXPERIENCE_INJECTION_POINT}}") # AI will write here
    
    doc.add_heading('EDUCATION', level=1)
    doc.add_paragraph("Bachelor in Business Management - MIS | 2021")
    
    doc.add_heading('SKILLS', level=1)
    doc.add_paragraph("Python, SQL, AWS, Docker, Kubernetes, TensorFlow, PyTorch, Generative AI")

    doc.save("Master_Resume_For_AI.docx")
    print("✅ 'Master_Resume_For_AI.docx' created. Please open it and PASTE your real details, but KEEP the {{BRACKETS}}.")

if __name__ == "__main__":
    create_master()