import os
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
import google.generativeai as genai
from pathlib import Path

# Configuration
CSV_FILE = "Sniper_Hourly_Leads.csv"
MASTER_RESUME = "Master_Resume.docx"
OUTPUT_DIR = "resumes"
GEMINI_API_KEY = "YOUR_GEMINI_API_KEY_HERE"  # Get from https://makersuite.google.com/app/apikeys

os.makedirs(OUTPUT_DIR, exist_ok=True)

# Configure Gemini
genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel('gemini-pro')

def generate_ai_content(job_title, company):
    """Use Gemini API to generate Summary + Project Description"""
    prompt = f"""
    I am applying for a {job_title} position at {company}.
    
    Generate TWO things:
    1. A 3-4 line PROFESSIONAL SUMMARY tailored to this role (highlight Python, ML, AWS, MLOps, IoT)
    2. A 5-6 line RELEVANT PROJECT DESCRIPTION that shows why I'm a perfect fit
    
    Format as:
    SUMMARY:
    [Your summary here]
    
    PROJECT:
    [Your project here]
    """
    
    try:
        response = model.generate_content(prompt)
        text = response.text
        
        # Parse response
        if "SUMMARY:" in text and "PROJECT:" in text:
            summary = text.split("SUMMARY:")[1].split("PROJECT:")[0].strip()
            project = text.split("PROJECT:")[1].strip()
            return summary, project
        else:
            return text[:200], text[200:400]
    except Exception as e:
        print(f"   ⚠️ Gemini API error: {e}")
        return "Experienced professional applying for this role.", "Relevant project experience in the domain."

def create_tailored_resume(job_title, company, summary, project):
    """Create tailored Word doc from Master_Resume template"""
    
    try:
        doc = Document(MASTER_RESUME)
    except FileNotFoundError:
        print(f"❌ {MASTER_RESUME} not found. Create it first!")
        return None
    
    # Replace placeholders
    for paragraph in doc.paragraphs:
        if "{{DYNAMIC_SUMMARY_INJECTION}}" in paragraph.text:
            paragraph.text = summary
        elif "{{DYNAMIC_PROJECT_INJECTION}}" in paragraph.text:
            paragraph.text = project
    
    # Safe filename
    safe_company = "".join(c if c.isalnum() or c in (" ", "-", "_") else "" for c in str(company)).strip().replace(" ", "_")[:30]
    safe_title = "".join(c if c.isalnum() or c in (" ", "-", "_") else "" for c in str(job_title)).strip().replace(" ", "_")[:30]
    
    filename = f"Manidhar_{safe_company}_{safe_title}.docx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    
    try:
        doc.save(filepath)
        return filepath
    except Exception as e:
        print(f"   ⚠️ Failed to save {filename}: {e}")
        return None

def main():
    # Read CSV
    if not os.path.exists(CSV_FILE):
        print(f"❌ {CSV_FILE} not found. Run sniper_bot_final.py first!")
        return
    
    df = pd.read_csv(CSV_FILE)
    
    # CRITICAL: Clean NaN values
    print("🧹 Cleaning CSV (removing NaN rows)...")
    df = df.dropna(subset=['Job Title', 'Company'])
    df = df[df['Job Title'].str.strip() != '']
    df = df[df['Company'].str.strip() != '']
    
    if df.empty:
        print("❌ No valid jobs in CSV after cleanup!")
        return
    
    print(f"✅ Found {len(df)} valid jobs to process.\n")
    
    created_resumes = []
    
    for i, row in df.iterrows():
        job_title = row['Job Title'].strip()
        company = row['Company'].strip()
        apply_link = row.get('Apply Link', '')
        
        print(f"[{i+1}/{len(df)}] Processing: {company} | {job_title}")
        
        # Generate AI content
        print(f"   🤖 Calling Gemini API...")
        summary, project = generate_ai_content(job_title, company)
        
        # Create tailored resume
        print(f"   📄 Creating tailored resume...")
        filepath = create_tailored_resume(job_title, company, summary, project)
        
        if filepath:
            created_resumes.append({
                "Job Title": job_title,
                "Company": company,
                "Apply Link": apply_link,
                "Resume File": filepath
            })
            print(f"   ✅ Created: {os.path.basename(filepath)}\n")
        else:
            print(f"   ❌ Failed to create resume.\n")
    
    # Generate Mission_Control.html
    if created_resumes:
        print(f"\n📋 Generating Mission_Control.html with {len(created_resumes)} jobs...")
        generate_html_tracker(created_resumes)
        print("✅ All done! Open 'Mission_Control.html' on your phone.")
    else:
        print("❌ No resumes created.")

def generate_html_tracker(resumes):
    """Create mobile-friendly HTML tracker"""
    html = [
        "<!DOCTYPE html>",
        "<html>",
        "<head>",
        "<meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width, initial-scale=1'>",
        "<title>🚀 Mission Control - AIM1000</title>",
        "<style>",
        "body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif; padding: 12px; background: #0a1428; color: #e0e0e0; }",
        "h1 { color: #00d4ff; text-align: center; margin-bottom: 5px; }",
        ".subtitle { text-align: center; color: #888; margin-bottom: 20px; font-size: 12px; }",
        ".job-card { background: #1a2332; border-left: 4px solid #00d4ff; padding: 12px; margin: 10px 0; border-radius: 4px; }",
        ".company { font-weight: bold; color: #00d4ff; font-size: 14px; }",
        ".role { color: #ccc; font-size: 13px; margin: 4px 0; }",
        ".buttons { display: flex; gap: 8px; margin-top: 10px; }",
        "a { display: inline-block; padding: 8px 12px; border-radius: 4px; text-decoration: none; font-size: 12px; font-weight: bold; }",
        ".btn-apply { background: #00d4ff; color: #000; flex: 1; text-align: center; }",
        ".btn-apply:hover { background: #00f0ff; }",
        ".btn-download { background: #28a745; color: #fff; flex: 1; text-align: center; }",
        ".btn-download:hover { background: #34c759; }",
        ".counter { text-align: center; color: #888; margin-top: 20px; font-size: 12px; }",
        "</style>",
        "</head>",
        "<body>",
        "<h1>🚀 MISSION CONTROL</h1>",
        f"<p class='subtitle'>AIM1000 - Auto Job Application | {len(resumes)} Jobs Ready to Apply</p>",
    ]
    
    for idx, resume in enumerate(resumes, 1):
        resume_rel_path = os.path.relpath(resume['Resume File'], '.').replace('\\', '/')
        html.extend([
            f"<div class='job-card'>",
            f"<div class='counter'>{idx}/{len(resumes)}</div>",
            f"<div class='company'>{resume['Company']}</div>",
            f"<div class='role'>{resume['Job Title']}</div>",
            f"<div class='buttons'>",
            f"<a href='{resume['Apply Link']}' target='_blank' class='btn-apply'>📲 APPLY</a>",
            f"<a href='{resume_rel_path}' download class='btn-download'>📥 CV</a>",
            f"</div>",
            f"</div>",
        ])
    
    html.extend([
        "<div class='counter'>💡 Tip: Share this HTML file to your phone. Tap APPLY to open job page, then download & upload CV.</div>",
        "</body>",
        "</html>"
    ])
    
    with open("Mission_Control.html", "w", encoding="utf-8") as f:
        f.write("\n".join(html))
    
    # Open in browser (Windows)
    try:
        os.startfile(os.path.abspath("Mission_Control.html"))
    except:
        pass

if __name__ == "__main__":
    main()